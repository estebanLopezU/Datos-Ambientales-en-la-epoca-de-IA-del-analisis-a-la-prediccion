# -*- coding: utf-8 -*-
"""
Validacion final de los libros entregados.

Comprueba, con lectura de bajo consumo (read_only + XML):
  * integridad del .xlsx y lista de hojas
  * ninguna hoja excede el limite de filas de Excel
  * encabezados, fecha real de Excel y formatos numericos
  * cuadre de filas por hoja contra el resumen anual
  * resumen anual y hoja de calidad
"""
import sys
import zipfile
import openpyxl

LIMITE = 1_048_576
VARS = ["precipitacion", "temperatura", "humedad_relativa", "nivel", "caudal",
        "velocidad_viento", "direccion_viento", "presion_barometrica",
        "evapotranspiracion", "radiacion_solar"]

fallos = []


def chequear(cond, msg):
    print(("  OK   " if cond else "  FALLA") + " " + msg)
    if not cond:
        fallos.append(msg)


def validar(ruta, total_esperado):
    print("=" * 74)
    print(f"VALIDANDO: {ruta}")
    print("=" * 74)

    # 1. Integridad del ZIP/xlsx
    with zipfile.ZipFile(ruta) as z:
        malo = z.testzip()
        chequear(malo is None, f"archivo .xlsx integro (zip valido): {malo}")
        nombres = z.namelist()
    chequear(any(n.startswith("xl/worksheets/sheet") for n in nombres),
             "contiene hojas de trabajo")

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    hojas = wb.sheetnames
    print(f"\n  hojas ({len(hojas)}): {hojas}\n")

    chequear(hojas[0] == "Portada", "la primera hoja es Portada")
    chequear("Resumen_Anual" in hojas, "existe hoja Resumen_Anual")
    chequear("Calidad_Datos" in hojas, "existe hoja Calidad_Datos")
    anios = [h for h in hojas if h.isdigit()]
    chequear(len(anios) >= 20, f"hay {len(anios)} hojas anuales")
    chequear([int(a) for a in anios] == sorted(int(a) for a in anios),
             "las hojas anuales estan en orden cronologico")

    # 2. Limite de filas y cuadre de totales
    print()
    filas_por_anio = {}
    for h in anios:
        ws = wb[h]
        filas = (ws.max_row or 0) - 1
        filas_por_anio[int(h)] = filas
        if filas + 1 > LIMITE:
            chequear(False, f"hoja {h} excede el limite de Excel ({filas:,})")
    chequear(all(f + 1 <= LIMITE for f in filas_por_anio.values()),
             f"ninguna hoja excede 1.048.576 filas "
             f"(max={max(filas_por_anio.values()):,})")
    total_hojas = sum(filas_por_anio.values())
    chequear(total_hojas == total_esperado,
             f"total filas en hojas = {total_hojas:,} (esperado {total_esperado:,})")

    # 3. Contenido de una hoja de datos
    print()
    ws = wb[anios[5]]
    enc = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    esperado = ["Fecha y hora"] + [
        {"precipitacion": "Precipitacion (mm)", "temperatura": "Temperatura (C)",
         "humedad_relativa": "Humedad rel. (%)", "nivel": "Nivel",
         "caudal": "Caudal", "velocidad_viento": "Vel. viento (m/s)",
         "direccion_viento": "Dir. viento (grados)",
         "presion_barometrica": "Presion (hPa)",
         "evapotranspiracion": "Evapotranspiracion (mm)",
         "radiacion_solar": "Radiacion solar (W/m2)"}[v] for v in VARS]
    chequear(enc == esperado, f"encabezados correctos en hoja {anios[5]}")
    chequear("brillo" not in " ".join(map(str, enc)).lower(),
             "columna 'brillo' (vacia) excluida")

    fila2 = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    chequear(fila2[0].__class__.__name__ == "datetime",
             f"la fecha es datetime de Excel: {fila2[0]!r}")
    chequear(isinstance(fila2[2], (int, float)),
             f"la temperatura es numerica: {fila2[2]!r}")

    # 4. Resumen anual
    print()
    ws = wb["Resumen_Anual"]
    filas = list(ws.iter_rows(values_only=True))
    chequear(filas[0][0] == "Ano", "encabezado del resumen anual")
    n_resumen = len(filas) - 1
    chequear(n_resumen == len(anios),
             f"resumen cubre los {len(anios)} anos ({n_resumen})")
    suma_reg = sum(r[1] for r in filas[1:] if isinstance(r[1], (int, float)))
    chequear(suma_reg == total_esperado,
             f"suma de registros del resumen = {suma_reg:,}")

    cuadran = all(filas_por_anio.get(r[0]) == r[1] for r in filas[1:])
    chequear(cuadran, "registros del resumen cuadran con las hojas por ano")
    dias_ok = all(isinstance(r[2], (int, float)) and 0 < r[2] <= 366
                  for r in filas[1:])
    chequear(dias_ok, "dias con datos en rango plausible (1-366)")
    temp_ok = all(isinstance(r[3], (int, float)) for r in filas[1:])
    chequear(temp_ok, "temperatura media numerica en todos los anos")

    # 5. Calidad
    print()
    ws = wb["Calidad_Datos"]
    filas = list(ws.iter_rows(values_only=True))
    chequear(filas[0][0] == "Variable", "encabezado de calidad de datos")
    vars_cal = [r[0] for r in filas[1:]]
    chequear(any("brillo" in str(v).lower() for v in vars_cal),
             "'brillo' documentado en la hoja de calidad")
    fuera = {r[0]: r[5] for r in filas[1:]}
    print(f"  valores fuera de rango: {fuera}")
    chequear(any(isinstance(v, (int, float)) and v > 0 for v in fuera.values()),
             "se detectan valores centinela/errores documentados")

    wb.close()
    print()
    return len(hojas), total_hojas


if __name__ == "__main__":
    CASOS = [
        (r"C:\Users\esteb\OneDrive\Desktop\Est Campus Unal\Excel_Organizado"
         r"\Estacion_24_La_Nubia.xlsx", 2_256_189),
        (r"C:\Users\esteb\OneDrive\Desktop\Est Campus Unal\Excel_Organizado"
         r"\Estacion_29_Palogrande_Postgrados.xlsx", 2_544_220),
    ]
    for ruta, tot in CASOS:
        n_hojas, n_filas = validar(ruta, tot)
        print(f"  >> {n_hojas} hojas, {n_filas:,} filas\n")

    print("=" * 74)
    if fallos:
        print(f"RESULTADO: {len(fallos)} FALLA(S)")
        for f in fallos:
            print("  - " + f)
        sys.exit(1)
    print("RESULTADO: todas las validaciones pasaron")
